"""Guruhga kun yakuni: bugun bajarilganlar Excel'da, faqat shu guruhniki, muddatidami belgisi bilan."""
import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from puls import config, db, report
from puls.ingest import Incoming, ingest

CHAT, OTHER = -100321, -100654


@pytest.fixture()
def conn(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "t.db")
    db.init()
    c = db.connect()
    c.execute("INSERT INTO people(id, name, username) VALUES(2, 'Ulug''bek Nazarov', 'demo_ulugbek')")
    yield c
    c.close()


def say(conn, chat, mid, text, when, reply=None, uid=1, name="Javohir Jumaboyev"):
    ingest(conn, Incoming(chat, "G", mid, uid, name, "javohir_fs", when, text, reply), use_ai=False)


def test_excel_has_only_today_done_tasks_of_this_chat(conn):
    say(conn, CHAT, 1, "#task #Portal Login sahifasi @demo_ulugbek 20-oktyabrgacha", datetime(2026, 10, 15, 9, 0))
    say(conn, CHAT, 2, "#task #Portal Bosh sahifa @demo_ulugbek", datetime(2026, 10, 20, 9, 0))
    say(conn, CHAT, 3, "#task #Portal Hisobot", datetime(2026, 10, 20, 9, 0))
    say(conn, OTHER, 1, "#task #Arxiv Skaner", datetime(2026, 10, 20, 9, 0))
    say(conn, CHAT, 4, "#done", datetime(2026, 10, 21, 10, 5), reply=1, uid=2, name="Ulug'bek Nazarov")
    say(conn, CHAT, 5, "#done", datetime(2026, 10, 21, 16, 40), reply=2, uid=2, name="Ulug'bek Nazarov")
    say(conn, OTHER, 2, "#done", datetime(2026, 10, 21, 11, 0), reply=1)
    data, name, caption = report.build(conn, CHAT, date(2026, 10, 21))
    ws = load_workbook(io.BytesIO(data)).active
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True)]
    assert [(r[1], r[3], r[6], r[9]) for r in rows] == [("Login sahifasi", "Ulug'bek Nazarov", "10:05", "yo'q"),
                                                        ("Bosh sahifa", "Ulug'bek Nazarov", "16:40", None)]
    assert len(load_workbook(io.BytesIO(data)).sheetnames) == 1
    assert name == "Bajarilgan_vazifalar_21.10.2026.xlsx"
    assert "2 ta vazifa bajarildi" in caption and "Ulug'bek (2)" in caption and "Ochiq qoldi: 1" in caption


def test_nothing_done_gives_none(conn):
    say(conn, CHAT, 1, "#task #Portal Login sahifasi", datetime(2026, 10, 20, 9, 0))
    assert report.build(conn, CHAT, date(2026, 10, 21)) is None
    assert report.chats(conn) == [CHAT]


def test_telegram_decoration_removed_and_rename_follows_everywhere(conn):
    from puls import db as _db
    say(conn, CHAT, 1, "#task #Portal Login sahifasi", datetime(2026, 10, 21, 9, 0), uid=7, name="Javohir ⚜️")
    assert conn.execute("SELECT assignee FROM tasks").fetchone()["assignee"] == "Javohir"
    _db.rename_person(conn, "Javohir", "Javohir Jumaboyev")
    conn.execute("UPDATE people SET name_locked = 1 WHERE id = 7")
    say(conn, CHAT, 2, "#done", datetime(2026, 10, 21, 15, 0), reply=1, uid=7, name="Javohir ⚜️")
    assert conn.execute("SELECT name FROM people WHERE id = 7").fetchone()["name"] == "Javohir Jumaboyev"
    row = conn.execute("SELECT assignee, status FROM tasks").fetchone()
    assert tuple(row) == ("Javohir Jumaboyev", "done")
    _, _, caption = report.build(conn, CHAT, date(2026, 10, 21))
    assert "Javohir (1)" in caption
