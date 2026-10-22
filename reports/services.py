"""Guruhga kun yakuni: shu guruhda bugun bajarilgan vazifalar Excel fayl bo'lib, qisqa xulosa izohda."""
import io
from collections import Counter
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analytics.services import uz_date
from projects.models import Chat, Task

HEAD = ["№", "Vazifa", "Loyiha", "Bajardi", "Bergan", "Berilgan", "Bajarilgan", "Necha kunda", "Muddat", "Muddatida"]
WIDTH = [5, 46, 26, 24, 22, 12, 12, 14, 12, 13]


def _dm(ts: str) -> str:
    return "%s.%s.%s" % (ts[8:10], ts[5:7], ts[:4])


def safe(value):
    """Excel qiymatni formula deb o'qimasligi uchun: = + - @ bilan boshlangan matn oldiga ' qo'yiladi."""
    return "'" + value if isinstance(value, str) and value[:1] in ("=", "+", "-", "@") else value


def done_today(chat_id: int, day: date) -> list[dict]:
    start = datetime.combine(day, time.min)
    rows = Task.objects.filter(project__chat_id=chat_id, status="done", done_at__gte=start, done_at__lt=start + timedelta(days=1)).order_by("done_at")
    return [{"title": t.title, "project": t.project.name, "assignee": t.assignee.name if t.assignee else None,
             "created_at": t.created_at.strftime("%Y-%m-%d %H:%M:%S"), "done_at": t.done_at.strftime("%Y-%m-%d %H:%M:%S"),
             "deadline": t.deadline.isoformat() if t.deadline else None,
             "giver": t.created_msg.user.name if t.created_msg and t.created_msg.user else None}
            for t in rows.select_related("project", "assignee", "created_msg__user")]


def build(chat_id: int, day: date) -> tuple[bytes, str, str] | None:
    """(xlsx baytlari, fayl nomi, izoh). Bugun hech narsa bajarilmagan bo'lsa None."""
    rows = done_today(chat_id, day)
    if not rows:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Bajarilgan"
    ws.append(["Bugun bajarilgan vazifalar, %s" % uz_date(day)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(HEAD)
    line = Side(style="thin", color="D1D1D6")
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4FA3")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, t in enumerate(rows, 1):
        days = (date.fromisoformat(t["done_at"][:10]) - date.fromisoformat(t["created_at"][:10])).days
        on_time = "" if not t["deadline"] else "ha" if t["done_at"][:10] <= t["deadline"] else "yo'q"
        ws.append([i, safe(t["title"]), safe(t["project"]), t["assignee"] or "", t["giver"] if t["giver"] and t["giver"] != t["assignee"] else "",
                   _dm(t["created_at"]), t["done_at"][11:16], days, _dm(t["deadline"]) if t["deadline"] else "", on_time])
        if on_time == "yo'q":
            ws.cell(ws.max_row, 10).font = Font(color="B42318", bold=True)
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for c in r:
            c.border = Border(bottom=line)
            if c.row > 3:
                c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 3))
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(HEAD)), ws.max_row)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)

    who = Counter(t["assignee"] for t in rows if t["assignee"]).most_common(3)
    open_qs = Task.objects.filter(project__chat_id=chat_id, status="open")
    open_n, late_n = open_qs.count(), open_qs.filter(deadline__lt=day).count()
    caption = "<b>Kun yakuni, %s</b>\nBugun %d ta vazifa bajarildi." % (uz_date(day), len(rows))
    if who:
        caption += "\nEng ko'p: %s." % ", ".join("%s (%d)" % (n.split(" ")[0], k) for n, k in who)
    caption += "\nOchiq qoldi: %d%s." % (open_n or 0, ", muddati o'tgan: %d" % late_n if late_n else "")
    return buf.getvalue(), "Bajarilgan_vazifalar_%s.xlsx" % day.strftime("%d.%m.%Y"), caption


def chats() -> list[int]:
    """Hisobot boradigan guruhlar: haqiqiy (namuna emas) va loyihasi bor guruhlar."""
    return list(Chat.objects.filter(demo=False, id__lt=0, projects__isnull=False).distinct().values_list("id", flat=True))


def panel_export(view: str, data: dict) -> tuple[bytes, str]:
    """Har ekran o'z jadvalini bitta toza varaqli alohida faylga chiqaradi. data — analytics.compute natijasi."""
    if view == "people":
        title, header, widths = "Xodimlar", ["Xodim", "Bo'lim", "Holat", "Fakt", "Ochiq", "Bajardi", "Bajarish % (30 kun)", "O'rtacha kun", "Baho"], [24, 13, 14, 34, 9, 10, 18, 13, 8]
        rows = [[p["name"], p["department"], p["state_label"], p["fact"], p["open"], p["done"], p["completion"], p["cycle"], p["score"]] for p in data["people"]]
    elif view == "tasks":
        title, header, widths = "Ochiq vazifalar", ["Xodim", "Loyiha", "Vazifa", "Necha kundan beri"], [24, 28, 56, 18]
        rows = [[p["name"], t["project"], t["title"], t["age"]] for p in data["people"] for t in p["open_tasks"]]
    elif view == "departments":
        title, header, widths = "Bo'limlar", ["Bo'lim", "Odamlar", "Ochiq", "Yangi", "Bajarildi", "Bajarilish %", "O'rtacha baho"], [16, 10, 9, 9, 11, 13, 14]
        rows = [[d["name"], d["count"], d["open"], d["created"], d["done"], d["pct"], d["score"]] for d in data["departments"]]
    else:
        title, header, widths = "Loyihalar", ["Loyiha", "Holat", "Asosiy sabab", "Progress %", "Ochiq", "Muddat", "Prognoz", "Yangi", "Bajarildi"], [28, 11, 52, 12, 9, 13, 13, 9, 11]
        rows = [[p["name"], p["status"], p["risks"][0]["text"] if p["risks"] else "", p["progress"], p["open_total"],
                 p["deadline_label"], p["eta_label"], p["created"], p["done"]] for p in data["projects"]]
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["Puls · %s · %s" % (title, data["period"]["range_label"])])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(header)
    for c in ws[3]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5"), Alignment(vertical="center")
    for i, r in enumerate(rows):
        ws.append([safe(v) for v in r])
        if i % 2:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="F5F6FA")
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "puls_%s_%s.xlsx" % (view, data["period"]["to"])
