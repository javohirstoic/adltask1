"""Guruhga kun yakuni: shu guruhda bugun bajarilgan vazifalar Excel fayl bo'lib, qisqa xulosa izohda."""
import io
import sqlite3
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analytics import uz_date

HEAD = ["№", "Vazifa", "Loyiha", "Bajardi", "Bergan", "Berilgan", "Bajarilgan", "Necha kunda", "Muddat", "Muddatida"]
WIDTH = [5, 46, 26, 24, 22, 12, 12, 14, 12, 13]


def _dm(ts: str) -> str:
    return "%s.%s.%s" % (ts[8:10], ts[5:7], ts[:4])


def done_today(conn: sqlite3.Connection, chat_id: int, day: date) -> list[dict]:
    return [dict(r) for r in conn.execute(
        "SELECT t.title, p.name AS project, t.assignee, t.created_at, t.done_at, t.deadline, pe.name AS giver "
        "FROM tasks t JOIN projects p ON p.id = t.project_id LEFT JOIN messages m ON m.id = t.created_msg "
        "LEFT JOIN people pe ON pe.id = m.user_id "
        "WHERE p.chat_id = ? AND t.status = 'done' AND substr(t.done_at, 1, 10) = ? ORDER BY t.done_at",
        (chat_id, day.isoformat()))]


def build(conn: sqlite3.Connection, chat_id: int, day: date) -> tuple[bytes, str, str] | None:
    """(xlsx baytlari, fayl nomi, izoh). Bugun hech narsa bajarilmagan bo'lsa None."""
    rows = done_today(conn, chat_id, day)
    if not rows:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Bajarilgan"
    ws.append(["Bugun bajarilgan vazifalar, %s" % uz_date(day)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(HEAD)
    line = Side(style="thin", color="D1D1D6")
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4FA3")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, t in enumerate(rows, 1):
        days = (date.fromisoformat(t["done_at"][:10]) - date.fromisoformat(t["created_at"][:10])).days
        on_time = "" if not t["deadline"] else "ha" if t["done_at"][:10] <= t["deadline"] else "yo'q"
        ws.append([i, t["title"], t["project"], t["assignee"] or "", t["giver"] if t["giver"] != t["assignee"] else "",
                   _dm(t["created_at"]), t["done_at"][11:16], days, _dm(t["deadline"]) if t["deadline"] else "", on_time])
        if on_time == "yo'q":
            ws.cell(ws.max_row, 10).font = Font(color="B42318", bold=True)
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for c in r:
            c.border = Border(bottom=line)
            if c.row > 3:
                c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 3))
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(HEAD)), ws.max_row)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)

    who = Counter(t["assignee"] for t in rows if t["assignee"]).most_common(3)
    open_n, late_n = conn.execute(
        "SELECT count(*), sum(t.deadline IS NOT NULL AND t.deadline < ?) FROM tasks t JOIN projects p ON p.id = t.project_id "
        "WHERE p.chat_id = ? AND t.status = 'open'", (day.isoformat(), chat_id)).fetchone()
    caption = "<b>Kun yakuni, %s</b>\nBugun %d ta vazifa bajarildi." % (uz_date(day), len(rows))
    if who:
        caption += "\nEng ko'p: %s." % ", ".join("%s (%d)" % (n.split(" ")[0], k) for n, k in who)
    caption += "\nOchiq qoldi: %d%s." % (open_n or 0, ", muddati o'tgan: %d" % late_n if late_n else "")
    return buf.getvalue(), "Bajarilgan_vazifalar_%s.xlsx" % day.strftime("%d.%m.%Y"), caption


def chats(conn: sqlite3.Connection) -> list[int]:
    """Hisobot boradigan guruhlar: haqiqiy (namuna emas) va loyihasi bor guruhlar."""
    return [r["id"] for r in conn.execute(
        "SELECT DISTINCT c.id FROM chats c JOIN projects p ON p.chat_id = c.id WHERE c.demo = 0 AND c.id < 0")]
